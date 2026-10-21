#!/usr/bin/env python3
"""
docusign_to_flexera.py  -  DocuSign envelope usage -> Excel -> Flexera subscription custom fields.

One run does all of this:
  1. DOCUSIGN   Authenticate (tokens kept in .env and updated there automatically), count envelopes
                per month incrementally, and write the usage workbook (same as docusign_monthly.py):
                  sheet "MonthlyUsage"  one row per month, split by status
                  sheet "TotalsToDate"  single-row rollup
  2. CALCULATE  Consumed  = envelopes to date (TotalsToDate)
                Remaining = DS_TOTAL_ENVELOPES - Consumed
                Both are written back into .env (DS_ENVELOPES_CONSUMED / DS_ENVELOPES_REMAINING).
  3. FLEXERA    Set the subscription custom fields on "Docusign Envelope":
                  Total Envelops / Envelops Consumed / Envelops Remaining
                Every other custom field on the subscription is kept (read first, merge, write back).

TOKENS LIVE IN .env
  DS_ACCESS_TOKEN, DS_ACCESS_TOKEN_EXPIRES_UTC and DS_REFRESH_TOKEN are read from .env.
  When a new access token is generated - and when DocuSign rotates the refresh token - the new
  values are written straight back into .env (other lines and comments are left untouched).
  Tokens are never printed or written to the log.
  A stored access token is reused until 5 minutes before it expires; an expired or rejected one is
  replaced automatically using the refresh token.

.env keys
  # DocuSign
  DS_ENV=production                    production | demo
  DS_INTEGRATION_KEY=...               (yml username)
  DS_CLIENT_SECRET=...                 (yml password)
  DS_REFRESH_TOKEN=...                 updated automatically when DocuSign rotates it
  DS_ACCESS_TOKEN=                     filled in automatically
  DS_ACCESS_TOKEN_EXPIRES_UTC=         filled in automatically
  DS_ACCOUNT_ID=...
  DS_START_DATE=2025-05-01
  DS_COUNT_MODE=fast                   fast | full
  DS_TZ_OFFSET=+00:00
  DS_FROM_TO_STATUS=created            created | changed
  DS_REBUILD=0                         1 = refetch every month
  DS_TIMEOUT=180
  DS_OUTPUT_FILE=docusign_envelope_usage.xlsx      (inside OUTPUT_DIR unless an absolute path)

  # Entitlement and results
  DS_TOTAL_ENVELOPES=339800            envelopes purchased
  DS_CONSUMED_BASIS=total              total = all envelopes (TotalsToDate.TotalEnvelopes)
                                       exclude_drafts = total minus status "Created" (unsent drafts)
  DS_ENVELOPES_CONSUMED=               written by the script after every run
  DS_ENVELOPES_REMAINING=              written by the script after every run
  DS_LAST_COUNTED_UTC=                 written by the script after every run

  # Flexera
  FLEXERA_TOKEN_URL / FLEXERA_CLIENT_ID / FLEXERA_CLIENT_SECRET
  FLEXERA_API_BASE=https://snowatlas-westeurope.flexera.eu
  DS_SUBSCRIPTION_ID=74aeb669-b626-4111-b00f-dd92f676386d
  DS_FIELD_TOTAL=Total Envelops
  DS_FIELD_CONSUMED=Envelops Consumed
  DS_FIELD_REMAINING=Envelops Remaining
  DS_DRY_RUN=yes                       yes = do not write to Flexera. No = write.

  # What to run, output, log
  RUN_STEPS=docusign,flexera           docusign | flexera | docusign,flexera
                                       (flexera alone uses DS_ENVELOPES_CONSUMED from .env)
  OUTPUT_DIR=output
  LOG_MAX_MB=5                         log rotates to docusign_to_flexera_previous.log

Requires:  python -m pip install requests python-dotenv openpyxl
"""
import base64
import calendar
import json
import os
import sys
import time
import traceback
from datetime import datetime, timedelta, timezone
from pathlib import Path
from urllib.parse import urlparse

try:
    import requests
except ImportError:
    sys.exit("requests not installed ->  python -m pip install requests")
try:
    from dotenv import load_dotenv
except ImportError:
    sys.exit("python-dotenv not installed ->  python -m pip install python-dotenv")
try:
    from openpyxl import Workbook, load_workbook
except ImportError:
    sys.exit("openpyxl not installed ->  python -m pip install openpyxl")

BASE_DIR = Path(__file__).resolve().parent
ENV_PATH = Path(os.getenv("DOTENV_PATH") or BASE_DIR / ".env")
load_dotenv(ENV_PATH, override=True)

for _new, _old in (("FLEXERA_TOKEN_URL", "TOKEN_URL"), ("FLEXERA_CLIENT_ID", "CLIENT_ID"),
                   ("FLEXERA_CLIENT_SECRET", "CLIENT_SECRET"), ("FLEXERA_API_BASE", "API_BASE_URL")):
    if not (os.getenv(_new) or "").strip() and (os.getenv(_old) or "").strip():
        os.environ[_new] = os.environ[_old]


def env(key, default=""):
    return (os.getenv(key) or default).strip()


def required(key):
    val = env(key)
    if not val or val.startswith("PASTE_") or "<" in val:
        sys.exit(f"[config] {key} is missing in .env")
    return val


# ═════════════════════════════ CONFIG ═════════════════════════════
STEPS = {s.strip().lower() for s in env("RUN_STEPS", "docusign,flexera").split(",") if s.strip()}
if not STEPS or not STEPS <= {"docusign", "flexera"}:
    sys.exit("[config] RUN_STEPS must be docusign, flexera, or docusign,flexera")

_out = Path(env("OUTPUT_DIR", "output"))
OUTPUT_DIR = _out if _out.is_absolute() else BASE_DIR / _out
LOG_FILE = OUTPUT_DIR / "docusign_to_flexera.log"
LOG_MAX_BYTES = int(float(env("LOG_MAX_MB", "5")) * 1024 * 1024)

# DocuSign
ENVIRONMENT = env("DS_ENV", "production").lower()
AUTH_BASE = env("DS_AUTH_BASE") or ("https://account-d.docusign.com" if ENVIRONMENT == "demo"
                                    else "https://account.docusign.com")
START_DATE = env("DS_START_DATE", "2025-05-01")
_xlsx = Path(env("DS_OUTPUT_FILE", "docusign_envelope_usage.xlsx"))
OUTPUT_FILE = _xlsx if _xlsx.is_absolute() else OUTPUT_DIR / _xlsx
FROM_TO_STATUS = env("DS_FROM_TO_STATUS", "created")
COUNT_MODE = env("DS_COUNT_MODE", "fast").lower()
TZ_OFFSET = env("DS_TZ_OFFSET", "+00:00")
REBUILD = env("DS_REBUILD") == "1"
HTTP_TIMEOUT = int(env("DS_TIMEOUT", "180"))
ENVELOPES_PAGE_SIZE = 100
TOKEN_MARGIN = timedelta(minutes=5)

CONSUMED_BASIS = env("DS_CONSUMED_BASIS", "total").lower()
if CONSUMED_BASIS not in ("total", "exclude_drafts"):
    sys.exit("[config] DS_CONSUMED_BASIS must be total or exclude_drafts")

# Flexera
FIELDS_PATH = env("CF_FIELDS_PATH", "/api/saas/consolidated-view/v1/custom-fields")
VALUES_PATH = env("CF_VALUES_PATH", FIELDS_PATH.rstrip("/") + "/values")
SUBS_PATH = env("DS_SUBSCRIPTIONS_PATH", "/api/saas/consolidated-view/v1/subscriptions")
CATEGORY = env("DS_SAAS_CATEGORY", "License")
FIELD_TOTAL = env("DS_FIELD_TOTAL", "Total Envelops")
FIELD_CONSUMED = env("DS_FIELD_CONSUMED", "Envelops Consumed")
FIELD_REMAINING = env("DS_FIELD_REMAINING", "Envelops Remaining")
DRY_RUN = env("DS_DRY_RUN", "yes").lower() != "no"
MAX_RETRIES = 5

MONTHLY_SHEET = "MonthlyUsage"
TOTALS_SHEET = "TotalsToDate"
API_STATUSES = [
    ("completed", "Completed"), ("declined", "Declined"), ("voided", "Voided"), ("deleted", "Deleted"),
    ("timedout", "Expired"), ("correct", "Corrected"), ("created", "Created"), ("sent", "Sent"),
    ("delivered", "Delivered"), ("signed", "Signed"),
]
STATUS_COLUMNS = [col for _api, col in API_STATUSES]
DATE_COLUMNS = ["MonthStart", "MonthEnd", "Year", "MonthNumber", "MonthName", "MonthYear",
                "YearMonthKey", "DaysInMonth", "DaysCounted", "IsComplete"]
DERIVED_COLUMNS = ["InProgress", "TotalEnvelopes", "SumOfStatuses", "CountVariance"]
MONTHLY_HEADER = DATE_COLUMNS + STATUS_COLUMNS + DERIVED_COLUMNS + ["FetchedUTC"]


# ═════════════════════════════ LOG ═════════════════════════════
class _Tee:
    """Console + log file; lines ending in \\r (progress) go to the console only.
    The log rotates to <name>_previous.log when it passes LOG_MAX_MB."""
    def __init__(self, stream, path):
        self.stream, self.path, self.buf = stream, path, ""
        path.parent.mkdir(parents=True, exist_ok=True)
        if path.exists() and path.stat().st_size >= LOG_MAX_BYTES:
            self._rotate()
        self.f = open(path, "a", encoding="utf-8")

    def _rotate(self):
        os.replace(self.path, self.path.with_name(f"{self.path.stem}_previous{self.path.suffix}"))

    def write(self, text):
        self.stream.write(text)
        self.buf += text
        while True:
            cuts = [i for i in (self.buf.find("\n"), self.buf.find("\r")) if i >= 0]
            if not cuts:
                break
            cut = min(cuts)
            line, end, self.buf = self.buf[:cut], self.buf[cut], self.buf[cut + 1:]
            if end == "\n" and line.strip():
                self.f.write(f"{datetime.now():%Y-%m-%d %H:%M:%S}  {line}\n")
                self.f.flush()
                if self.f.tell() >= LOG_MAX_BYTES:
                    self.f.close()
                    self._rotate()
                    self.f = open(self.path, "a", encoding="utf-8")

    def flush(self):
        self.stream.flush()


def start_log():
    tee = _Tee(sys.__stdout__, LOG_FILE)
    tee.f.write("\n" + "=" * 70 + f"\n{datetime.now():%Y-%m-%d %H:%M:%S}  RUN START\n")
    tee.f.flush()
    sys.stdout = tee
    sys.stderr = tee


# ═════════════════════════════ .env WRITER ═════════════════════════════
def update_env_file(updates, label):
    """Set KEY=value in .env for every key in `updates` (every occurrence), appending missing keys.
    All other lines and comments are kept exactly. Written atomically."""
    updates = {k: "" if v is None else str(v) for k, v in updates.items()}
    lines = ENV_PATH.read_text(encoding="utf-8-sig").splitlines() if ENV_PATH.exists() else []
    out, done = [], set()
    for line in lines:
        s = line.strip()
        body = s[7:].lstrip() if s.startswith("export ") else s
        key = body.split("=", 1)[0].strip() if "=" in body and not s.startswith("#") else None
        if key in updates:
            out.append(f"{key}={updates[key]}")
            done.add(key)
        else:
            out.append(line)
    missing = [k for k in updates if k not in done]
    if missing:
        if out and out[-1].strip():
            out.append("")
        out += [f"{k}={updates[k]}" for k in missing]
    text = "\n".join(out) + "\n"
    tmp = ENV_PATH.with_name(ENV_PATH.name + ".tmp")
    try:
        tmp.write_text(text, encoding="utf-8")
        os.replace(tmp, ENV_PATH)
    except OSError:
        try:
            ENV_PATH.write_text(text, encoding="utf-8")
            if tmp.exists():
                tmp.unlink()
        except OSError as e:
            return False, str(e)
    os.environ.update(updates)
    print(f"[env] Updated .env: {label} ({', '.join(updates)})")
    return True, ""


# ═════════════════════════════ DOCUSIGN ═════════════════════════════
def parse_offset(text):
    text = text.strip()
    if not text or text.upper() == "Z":
        return timedelta(0)
    sign = -1 if text[0] == "-" else 1
    body = text.lstrip("+-").replace(":", "")
    if len(body) != 4:
        sys.exit(f"[config] DS_TZ_OFFSET '{text}' is not in +HH:MM form.")
    return sign * timedelta(hours=int(body[:2]), minutes=int(body[2:]))


OFFSET = parse_offset(TZ_OFFSET)


def to_utc_string(naive_local):
    return (naive_local - OFFSET).strftime("%Y-%m-%dT%H:%M:%SZ")


class DocuSign:
    def __init__(self):
        self.token = None
        self.base_path = None

    # ---- tokens ----
    def _stored_token_valid(self):
        tok, exp = env("DS_ACCESS_TOKEN"), env("DS_ACCESS_TOKEN_EXPIRES_UTC")
        if not tok:
            return False, "no access token in .env"
        if not exp:
            return True, "access token in .env (expiry unknown - will be checked)"
        try:
            exp_dt = datetime.strptime(exp, "%Y-%m-%dT%H:%M:%SZ").replace(tzinfo=timezone.utc)
        except ValueError:
            return True, "access token in .env (expiry unreadable - will be checked)"
        if datetime.now(timezone.utc) + TOKEN_MARGIN >= exp_dt:
            return False, f"access token in .env expired at {exp}"
        return True, f"access token from .env (valid until {exp})"

    def refresh(self, reason):
        key, secret, refresh = env("DS_INTEGRATION_KEY"), env("DS_CLIENT_SECRET"), env("DS_REFRESH_TOKEN")
        if not (key and secret and refresh):
            sys.exit(f"[docusign] {reason}, and DS_INTEGRATION_KEY / DS_CLIENT_SECRET / DS_REFRESH_TOKEN "
                     f"are not all set in .env - cannot get a new access token.")
        print(f"[docusign] Getting a new access token ({reason})")
        basic = base64.b64encode(f"{key}:{secret}".encode()).decode()
        resp = requests.post(f"{AUTH_BASE}/oauth/token", timeout=30,
                             headers={"Authorization": f"Basic {basic}"},
                             data={"grant_type": "refresh_token", "refresh_token": refresh})
        if resp.status_code != 200:
            sys.exit(f"[docusign] Token request failed ({resp.status_code}). Usual causes: the refresh token "
                     f"was already used/rotated elsewhere, it expired (~30 days unused), or the integration "
                     f"key/secret don't match. DocuSign said: {resp.text[:300]}")
        data = resp.json()
        expires = datetime.now(timezone.utc) + timedelta(seconds=int(data.get("expires_in") or 28800))
        updates = {"DS_ACCESS_TOKEN": data["access_token"],
                   "DS_ACCESS_TOKEN_EXPIRES_UTC": expires.strftime("%Y-%m-%dT%H:%M:%SZ")}
        rotated = data.get("refresh_token") and data["refresh_token"] != refresh
        if rotated:
            updates["DS_REFRESH_TOKEN"] = data["refresh_token"]
        ok, err = update_env_file(updates, "new access token" + (" + rotated refresh token" if rotated else ""))
        if not ok:
            # The rotated refresh token must not be lost - save it to a fixed fallback file
            fallback = OUTPUT_DIR / "docusign_tokens_NOT_SAVED.env"
            fallback.write_text("\n".join(f"{k}={v}" for k, v in updates.items()) + "\n", encoding="utf-8")
            print(f"[env] COULD NOT WRITE .env ({err}). The new tokens are in {fallback} - copy them into .env "
                  f"and delete that file, or the next run will fail.")
        self.token = data["access_token"]

    def authenticate(self):
        ok, why = self._stored_token_valid()
        if ok:
            print(f"[docusign] Using {why}")
            self.token = env("DS_ACCESS_TOKEN")
        else:
            self.refresh(why)
        resp = requests.get(f"{AUTH_BASE}/oauth/userinfo", timeout=30,
                            headers={"Authorization": f"Bearer {self.token}"})
        if resp.status_code == 401:
            self.refresh("stored access token was rejected")
            resp = requests.get(f"{AUTH_BASE}/oauth/userinfo", timeout=30,
                                headers={"Authorization": f"Bearer {self.token}"})
        if not resp.ok:
            sys.exit(f"[docusign] userinfo failed ({resp.status_code}): {resp.text[:300]}")
        accounts = resp.json()["accounts"]
        acct_id = env("DS_ACCOUNT_ID")
        if acct_id:
            account = next((a for a in accounts if a["account_id"] == acct_id), None)
            if account is None:
                sys.exit(f"[docusign] Account {acct_id} is not visible to this user.")
        else:
            account = next((a for a in accounts if a.get("is_default")), accounts[0])
        self.base_path = f"{account['base_uri']}/restapi/v2.1/accounts/{account['account_id']}"
        print(f"[docusign] Account: {account.get('account_name', '')} ({account['account_id']})")

    # ---- HTTP ----
    def get(self, url, params=None):
        refreshed = False
        for attempt in range(MAX_RETRIES):
            try:
                resp = requests.get(url, params=params, timeout=HTTP_TIMEOUT,
                                    headers={"Authorization": f"Bearer {self.token}", "Accept": "application/json"})
            except (requests.Timeout, requests.ConnectionError) as e:
                if attempt == MAX_RETRIES - 1:
                    sys.exit(f"[docusign] Request kept timing out ({e.__class__.__name__}). Try DS_TIMEOUT=300.")
                time.sleep(5 * (attempt + 1))
                continue
            if resp.status_code == 401 and not refreshed:
                self.refresh("access token expired during the run")
                refreshed = True
                continue
            if resp.status_code == 429:
                reset = resp.headers.get("X-RateLimit-Reset")
                wait = max(int(reset) - int(time.time()), 1) if reset else 2 ** attempt * 30
                print(f"    rate limited, sleeping {wait}s ...")
                time.sleep(min(wait, 3700))
                continue
            if resp.status_code >= 500:
                time.sleep(2 ** attempt)
                continue
            if 400 <= resp.status_code < 500:
                try:
                    err = resp.json()
                    detail = f"{err.get('errorCode', '')}: {err.get('message', resp.text)}"
                except ValueError:
                    detail = resp.text
                sys.exit(f"[docusign] Request rejected ({resp.status_code}): {detail[:300]}")
            left = resp.headers.get("X-RateLimit-Remaining")
            if left is not None and int(left) < 50:
                print(f"    warning: only {left} DocuSign API calls left this hour")
            return resp.json()
        raise RuntimeError(f"Giving up on {url}")

    # ---- counting ----
    def count_fast(self, ws, we):
        base = {"from_date": to_utc_string(ws), "to_date": to_utc_string(we),
                "from_to_status": FROM_TO_STATUS, "count": 1, "start_position": 0}
        counts = {}
        for api_status, column in API_STATUSES:
            data = self.get(f"{self.base_path}/envelopes", {**base, "status": api_status})
            counts[column] = int(data.get("totalSetSize", 0) or 0)
        data = self.get(f"{self.base_path}/envelopes", base)
        return counts, int(data.get("totalSetSize", 0) or 0)

    def count_full(self, ws, we):
        by_api = dict(API_STATUSES)
        counts, seen, unknown, pos = {c: 0 for c in STATUS_COLUMNS}, set(), {}, 0
        while True:
            data = self.get(f"{self.base_path}/envelopes", {
                "from_date": to_utc_string(ws), "to_date": to_utc_string(we),
                "from_to_status": FROM_TO_STATUS, "count": ENVELOPES_PAGE_SIZE,
                "start_position": pos, "order_by": "created", "order": "asc"})
            batch = data.get("envelopes", []) or []
            if not batch:
                break
            for e in batch:
                eid = e.get("envelopeId")
                if eid in seen:
                    continue
                seen.add(eid)
                raw = (e.get("status") or "").lower()
                col = by_api.get(raw)
                if col:
                    counts[col] += 1
                else:
                    unknown[raw] = unknown.get(raw, 0) + 1
            pos += len(batch)
            total = int(data.get("totalSetSize", 0) or 0)
            if total and pos >= total:
                break
        if unknown:
            print(f"    statuses seen but not modelled: {unknown}")
        return counts, len(seen)


def iter_months(start_date_str):
    start = datetime.strptime(start_date_str, "%Y-%m-%d")
    today = (datetime.now(timezone.utc) + OFFSET).replace(tzinfo=None)
    y, m = start.year, start.month
    while (y, m) <= (today.year, today.month):
        dim = calendar.monthrange(y, m)[1]
        current = (y, m) == (today.year, today.month)
        last = today.day if current else dim
        yield {"year": y, "month": m, "days_in_month": dim, "days_counted": last, "is_complete": not current,
               "window_start": datetime(y, m, 1), "window_end": datetime(y, m, last, 23, 59, 59)}
        m += 1
        if m > 12:
            m, y = 1, y + 1


def build_row(d, counts, total):
    y, m = d["year"], d["month"]
    in_prog = sum(counts.get(c, 0) for c in ("Created", "Sent", "Delivered", "Signed"))
    s = sum(counts.get(c, 0) for c in STATUS_COLUMNS)
    row = {"MonthStart": datetime(y, m, 1).date(), "MonthEnd": datetime(y, m, d["days_counted"]).date(),
           "Year": y, "MonthNumber": m, "MonthName": calendar.month_name[m],
           "MonthYear": f"{calendar.month_abbr[m]} {y}", "YearMonthKey": y * 100 + m,
           "DaysInMonth": d["days_in_month"], "DaysCounted": d["days_counted"], "IsComplete": d["is_complete"],
           "InProgress": in_prog, "TotalEnvelopes": total, "SumOfStatuses": s, "CountVariance": total - s,
           "FetchedUTC": datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M:%S")}
    row.update({c: counts.get(c, 0) for c in STATUS_COLUMNS})
    return row


def read_existing():
    if REBUILD or not OUTPUT_FILE.exists():
        return {}
    try:
        wb = load_workbook(OUTPUT_FILE)
    except Exception as e:
        print(f"[excel] Could not open {OUTPUT_FILE.name} ({e}) - starting a fresh workbook")
        return {}
    if MONTHLY_SHEET not in wb.sheetnames:
        return {}
    rows = list(wb[MONTHLY_SHEET].iter_rows(values_only=True))
    wb.close()
    if not rows:
        return {}
    header = [str(h) if h is not None else "" for h in rows[0]]
    out = {}
    for values in rows[1:]:
        rec = dict(zip(header, values))
        try:
            out[(int(rec["Year"]), int(rec["MonthNumber"]))] = rec
        except (KeyError, TypeError, ValueError):
            continue
    return out


def normalise(rec):
    out = {}
    for col in MONTHLY_HEADER:
        v = rec.get(col)
        if col in ("MonthStart", "MonthEnd") and isinstance(v, datetime):
            v = v.date()
        if col == "IsComplete":
            v = bool(v) and str(v).upper() != "FALSE"
        out[col] = v
    return out


def write_workbook(rows):
    rows = sorted(rows, key=lambda r: r["YearMonthKey"])
    wb = Workbook()
    ws = wb.active
    ws.title = MONTHLY_SHEET
    ws.append(MONTHLY_HEADER)
    for r in rows:
        ws.append([r.get(c) for c in MONTHLY_HEADER])
    for col in ("A", "B"):
        for cell in ws[col][1:]:
            cell.number_format = "yyyy-mm-dd"
    ws.freeze_panes = "A2"

    totals = wb.create_sheet(TOTALS_SHEET)
    complete = [r for r in rows if r.get("IsComplete")]
    header = (["LastUpdatedUTC", "RangeStart", "RangeEnd", "MonthsCovered", "CompleteMonths", "PartialMonths",
               "TotalEnvelopes", "TotalEnvelopesCompleteMonthsOnly", "AvgEnvelopesPerCompleteMonth"]
              + STATUS_COLUMNS + ["InProgress"])
    totals.append(header)
    grand = sum(int(r.get("TotalEnvelopes") or 0) for r in rows)
    comp_total = sum(int(r.get("TotalEnvelopes") or 0) for r in complete)
    rec = {"LastUpdatedUTC": datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M:%S"),
           "RangeStart": rows[0]["MonthStart"] if rows else None, "RangeEnd": rows[-1]["MonthEnd"] if rows else None,
           "MonthsCovered": len(rows), "CompleteMonths": len(complete), "PartialMonths": len(rows) - len(complete),
           "TotalEnvelopes": grand, "TotalEnvelopesCompleteMonthsOnly": comp_total,
           "AvgEnvelopesPerCompleteMonth": round(comp_total / len(complete), 1) if complete else 0,
           "InProgress": sum(int(r.get("InProgress") or 0) for r in rows)}
    for col in STATUS_COLUMNS:
        rec[col] = sum(int(r.get(col) or 0) for r in rows)
    totals.append([rec.get(c) for c in header])
    for col in ("B", "C"):
        for cell in totals[col][1:]:
            cell.number_format = "yyyy-mm-dd"

    OUTPUT_FILE.parent.mkdir(parents=True, exist_ok=True)
    tmp = OUTPUT_FILE.with_name(OUTPUT_FILE.stem + "_tmp" + OUTPUT_FILE.suffix)
    wb.save(tmp)
    try:
        os.replace(tmp, OUTPUT_FILE)
        print(f"[excel] Wrote {OUTPUT_FILE}")
    except PermissionError:
        print(f"[excel] {OUTPUT_FILE.name} is open (probably in Excel) - this run's workbook is in {tmp.name}. "
              f"Close Excel; the next run will rebuild it. Continuing with the counts.")
    return rec


def run_docusign():
    if COUNT_MODE not in ("fast", "full"):
        sys.exit("[config] DS_COUNT_MODE must be fast or full")
    ds = DocuSign()
    ds.authenticate()
    print(f"[docusign] Mode {COUNT_MODE} | from_to_status={FROM_TO_STATUS} | tz {TZ_OFFSET} | from {START_DATE}")
    existing = read_existing()
    if existing:
        print(f"[excel] {len(existing)} month(s) already in {OUTPUT_FILE.name}")
    rows, fetched, reused = [], 0, 0
    for d in iter_months(START_DATE):
        key = (d["year"], d["month"])
        label = f"{calendar.month_abbr[d['month']]} {d['year']}"
        prior = existing.get(key)
        if prior is not None and d["is_complete"]:
            if bool(prior.get("IsComplete")) and str(prior.get("IsComplete")).upper() != "FALSE":
                rows.append(normalise(prior))
                reused += 1
                continue
            print(f"  {label}: stored row was partial - refetching")
        why = "current month" if not d["is_complete"] else "new"
        print(f"  {label}: fetching ({why}, days 1-{d['days_counted']}) ...")
        counter = ds.count_fast if COUNT_MODE == "fast" else ds.count_full
        counts, total = counter(d["window_start"], d["window_end"])
        row = build_row(d, counts, total)
        rows.append(row)
        fetched += 1
        if row["CountVariance"]:
            print(f"    note: variance {row['CountVariance']:+d} "
                  f"(total {row['TotalEnvelopes']} vs statuses {row['SumOfStatuses']})")
    if not rows:
        sys.exit("[docusign] No months in range - check DS_START_DATE")
    summary = write_workbook(rows)
    print(f"[docusign] {len(rows)} months ({fetched} fetched, {reused} reused) | "
          f"envelopes to date: {summary['TotalEnvelopes']:,} (drafts/Created: {summary['Created']:,})")
    return summary


# ═════════════════════════════ FLEXERA ═════════════════════════════
class Flexera:
    def __init__(self):
        explicit = env("FLEXERA_API_BASE").rstrip("/")
        src = explicit if explicit and "<" not in explicit else required("FLEXERA_UPLOAD_URL")
        u = urlparse(src)
        self.base = f"{u.scheme}://{u.netloc}"
        self.s = requests.Session()
        self.s.headers["Accept"] = "application/json"
        self.refresh_token()
        print(f"[flexera] Authenticated ({self.base})")

    def refresh_token(self):
        resp = requests.post(required("FLEXERA_TOKEN_URL"), timeout=30,
                             headers={"Content-Type": "application/x-www-form-urlencoded"},
                             data={"grant_type": "client_credentials",
                                   "client_id": required("FLEXERA_CLIENT_ID"),
                                   "client_secret": required("FLEXERA_CLIENT_SECRET")})
        if not resp.ok:
            sys.exit(f"[flexera] Token request failed: HTTP {resp.status_code} {resp.text[:300]}")
        self.s.headers["Authorization"] = f"Bearer {resp.json()['access_token']}"

    def call(self, method, path, **kw):
        refreshed, resp = False, None
        for attempt in range(1, MAX_RETRIES + 1):
            resp = self.s.request(method, self.base + path, timeout=60, **kw)
            if resp.status_code == 401 and not refreshed:
                self.refresh_token()
                refreshed = True
                continue
            if resp.status_code in (429, 500, 502, 503, 504) and attempt < MAX_RETRIES:
                time.sleep(int(resp.headers.get("Retry-After") or 2 ** attempt))
                continue
            return resp
        return resp

    def paged(self, path, page_size):
        items, page = [], 1
        while True:
            resp = self.call("GET", path, params={"page_number": page, "page_size": page_size})
            if not resp.ok:
                sys.exit(f"[flexera] GET {path} failed: HTTP {resp.status_code} {resp.text[:200]}")
            body = resp.json()
            items.extend(body.get("items", []))
            if page >= ((body.get("pagination") or {}).get("total_pages") or 1):
                return items
            page += 1


def merge(existing, updates):
    upd = {k.lower(): (k, v) for k, v in updates.items()}
    out, done = [], set()
    for f in existing:
        if not isinstance(f, dict) or not f.get("name"):
            continue
        k = f["name"].strip().lower()
        if k in upd:
            out.append({"name": upd[k][0], "value": upd[k][1]})
            done.add(k)
        else:
            out.append({"name": f["name"], "value": f.get("value")})
    for k, (n, v) in upd.items():
        if k not in done:
            out.append({"name": n, "value": v})
    return out


def same(a, b):
    try:
        return float(a) == float(b)
    except (TypeError, ValueError):
        return str(a) == str(b)


def run_flexera(total, consumed, remaining):
    sub_id = required("DS_SUBSCRIPTION_ID")
    wanted = {FIELD_TOTAL: total, FIELD_CONSUMED: consumed, FIELD_REMAINING: remaining}
    print(f"[flexera] {'DRY RUN (nothing written)' if DRY_RUN else 'LIVE - values will be written'} | "
          + ", ".join(f"{k}={v:,}" for k, v in wanted.items()))
    fx = Flexera()

    defs = fx.paged(FIELDS_PATH, 100)
    lic = {(d.get("fieldName") or "").strip().lower(): d for d in defs
           if d.get("fieldName") and "license" in str(d.get("saasCategory") or "").lower()}
    missing = [n for n in wanted if n.lower() not in lic]
    if missing:
        sys.exit(f"[flexera] Subscription custom field(s) not found: {', '.join(missing)}. "
                 f"Existing: {', '.join(d['fieldName'] for d in lic.values()) or '(none)'}")
    updates = {}
    for n, v in wanted.items():
        d = lic[n.lower()]
        dtype = str(d.get("dataType") or "").lower()
        is_num = any(t in dtype for t in ("number", "int", "decimal", "double", "float", "currency"))
        updates[d["fieldName"]] = v if is_num else str(v)

    resp = fx.call("GET", f"{SUBS_PATH.rstrip('/')}/{sub_id}")
    existing = None
    if resp.ok:
        try:
            existing = resp.json().get("customFields")
        except ValueError:
            pass
    if not isinstance(existing, list):
        sys.exit(f"[flexera] Could not read the subscription's current custom fields (HTTP {resp.status_code}) "
                 f"- stopping WITHOUT writing so no other field can be wiped.")
    cur = {(f.get("name") or "").strip().lower(): f.get("value") for f in existing if isinstance(f, dict)}
    kept = [f"{f.get('name')}={f.get('value')}" for f in existing if isinstance(f, dict)
            and (f.get("name") or "").strip().lower() not in {k.lower() for k in updates}]
    print("[flexera] Current: " + ", ".join(f"{n}={cur.get(n.lower())!r}" for n in updates))
    print(f"[flexera] Other custom fields kept: {', '.join(kept) or '(none)'}")
    if all(same(cur.get(n.lower()), v) for n, v in updates.items()):
        print("[flexera] Already up to date - nothing to write")
        return
    payload = {"customFields": merge(existing, updates), "id": sub_id, "saasCategory": CATEGORY}
    if DRY_RUN:
        print(f"[flexera] Would send: {json.dumps(payload)}")
        return
    resp = fx.call("PATCH", VALUES_PATH, data=json.dumps(payload).encode("utf-8"),
                   headers={"Content-Type": "application/json"})
    if not resp.ok:
        print(f"[flexera] Body sent: {json.dumps(payload)}")
        sys.exit(f"[flexera] Update failed: HTTP {resp.status_code} {resp.text[:300]}")
    print("[flexera] Updated subscription: " + ", ".join(f"{n}={v:,}" if isinstance(v, (int, float)) else f"{n}={v}"
                                                         for n, v in updates.items()))


# ═════════════════════════════ MAIN ═════════════════════════════
def to_int(key):
    raw = env(key).replace(",", "")
    if not raw:
        return None
    try:
        return int(float(raw))
    except ValueError:
        sys.exit(f"[config] {key}='{env(key)}' is not a number")


def main():
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    total = to_int("DS_TOTAL_ENVELOPES")
    if total is None:
        sys.exit("[config] Set DS_TOTAL_ENVELOPES (envelopes purchased) in .env, e.g. DS_TOTAL_ENVELOPES=339800")
    print(f"[config] Steps: {', '.join(s for s in ('docusign', 'flexera') if s in STEPS)} | "
          f"total purchased: {total:,} | output: {OUTPUT_DIR}")

    if "docusign" in STEPS:
        summary = run_docusign()
        consumed = int(summary["TotalEnvelopes"])
        if CONSUMED_BASIS == "exclude_drafts":
            consumed -= int(summary.get("Created") or 0)
        remaining = total - consumed
        print(f"[result] Consumed {consumed:,} ({CONSUMED_BASIS}) | Remaining {remaining:,} of {total:,}")
        if remaining < 0:
            print(f"[warn] Consumed is more than the purchased total by {-remaining:,}")
        update_env_file({"DS_ENVELOPES_CONSUMED": consumed, "DS_ENVELOPES_REMAINING": remaining,
                         "DS_LAST_COUNTED_UTC": datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")},
                        "envelope counts")
    else:
        consumed = to_int("DS_ENVELOPES_CONSUMED")
        if consumed is None:
            sys.exit("[config] RUN_STEPS=flexera needs DS_ENVELOPES_CONSUMED in .env (run the docusign step first)")
        remaining = total - consumed
        print(f"[result] From .env (counted {env('DS_LAST_COUNTED_UTC') or 'unknown'}): "
              f"Consumed {consumed:,} | Remaining {remaining:,}")

    if "flexera" in STEPS:
        run_flexera(total, consumed, remaining)


if __name__ == "__main__":
    start_log()
    try:
        main()
        print("[run] Finished OK")
    except SystemExit as e:
        if e.code not in (None, 0):
            if not isinstance(e.code, int):
                print(e.code)
            print("[run] Stopped")
            sys.exit(1)
        print("[run] Finished OK")
    except KeyboardInterrupt:
        print("\n[run] Interrupted")
        sys.exit(1)
    except Exception:
        traceback.print_exc()
        print("[run] Crashed")
        sys.exit(1)
